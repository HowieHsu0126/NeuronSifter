"""Materialize audited public A/B/C evidence into the retrospective registry."""

from __future__ import annotations

import argparse
import csv
from datetime import datetime, timezone
import hashlib
import json
import math
from pathlib import Path
import sys
from typing import Any, Iterable, Mapping

import numpy as np
from jsonschema import Draft202012Validator
import openpyxl
from PIL import Image
import yaml

REPO_ROOT = Path(__file__).resolve().parents[3]
if str(REPO_ROOT) not in sys.path:
    sys.path.insert(0, str(REPO_ROOT))

from Libs.scripts.builders.build_retrospective_replay_benchmark import (  # noqa: E402
    EVIDENCE_FIELDS,
    FORMAL_GRADES,
    qualify_candidate_sets,
    validate_evidence_rows,
)
from Libs.utils.evaluation.retrospective_statistics import build_summary_arm  # noqa: E402


CONTRACT_PATH = REPO_ROOT / "Libs/configs/data/retrospective_public_therapeutic_evidence.yaml"
CONTRACT_SCHEMA_PATH = REPO_ROOT / "Libs/configs/schema/retrospective_public_therapeutic_evidence.schema.json"
MATERIALIZATION_SCHEMA_PATH = REPO_ROOT / "Libs/configs/schema/retrospective_public_evidence_materialization.schema.json"
BUILDER_PATH = REPO_ROOT / "Libs/scripts/builders/build_retrospective_public_evidence_materialization.py"
MANIFEST_PATH = REPO_ROOT / "Input/manifests/retrospective_public/evidence_materialization.tsv"
EVIDENCE_OUTPUT_PATH = REPO_ROOT / "Input/manifests/retrospective_public/evidence_registry.tsv"
REJECTION_OUTPUT_PATH = REPO_ROOT / "Output/audits/retrospective_public_evidence_rejection_catalog.tsv"
AUDIT_OUTPUT_PATH = REPO_ROOT / "Output/audits/retrospective_public_evidence_materialization.json"

MANIFEST_FIELDS = (
    "record_uid", "record_status", "evidence_uid", "publication_uid", "study_uid",
    "publication_date", "source_snapshot_uid", "source_artifact_path", "source_locator",
    "source_content_hash", "extraction_artifact_hash", "extraction_method", "extraction_qc_status",
    "public_access_status", "measurement_protocol", "biological_context_uid", "common_control_uid",
    "structured_action_uid", "action_family_uid", "target_program_uid", "regimen_family_uid",
    "objective_uid", "candidate_set_uid", "outcome_uid", "arm_mean", "biological_replicate_count",
    "uncertainty", "uncertainty_kind", "unit", "outcome_visibility", "target_retrieval_blocked",
    "tool_output_only", "evidence_role", "evidence_grade", "decision_domain_classification",
    "split_role", "rejection_reason",
)
REJECTION_FIELDS = (
    "record_uid", "publication_uid", "study_uid", "publication_date", "source_artifact_path",
    "source_locator", "source_content_hash", "evidence_grade", "extraction_qc_status", "rejection_reason",
)
SUMMARY_KINDS = {"extracted_se", "source_reported_sem"}
SD_KINDS = {"extracted_sd", "source_reported_sd"}
BOOTSTRAP_RESAMPLES = 2000
BOOTSTRAP_SEED = 3407
MIN_INDEPENDENT_GROUPS_FOR_REPLAY_ROLES = 3
FIGURE3F_AXIS_ZERO_PX = 229
FIGURE3F_AXIS_TOP_PX = 56
FIGURE3F_TICK_ROWS = (229, 186, 143, 100, 56)
FIGURE3F_PANEL_Y_OFFSET = 610
FIGURE3F_PANEL_X_OFFSET = 500
FIGURE3F_BAR_TOP_ROWS = {
    "DMSO": 188,
    "CompoundB": 179,
    "CompoundD": 165,
    "CompoundE": 138,
    "CompoundG": 188,
}
FIGURE3F_BAR_CENTERS = {"DMSO": 82, "CompoundB": 114, "CompoundD": 145, "CompoundE": 177, "CompoundG": 208}
ELIFE109518_FIG5_AXIS_TICK_ROWS = (2334, 2229, 2125, 2020)
ELIFE109518_FIG5_AXIS_TICK_VALUES = (0.0, 50.0, 100.0, 150.0)
ELIFE109518_FIG5_DOT_CENTERS = {
    "WT-75.1:untreated": ((315, 2303), (344, 2311), (329, 2314), (302, 2314)),
    "GD2-1260:untreated": ((409, 2209), (409, 2246), (409, 2271)),
    "GD2-1260:SapC-DOPS": ((476, 2263), (517, 2266), (497, 2273), (497, 2292)),
    "GD2-1260:SapC-DOPS-fGCase": ((584, 2314), (605, 2315), (563, 2322)),
}
ELIFE69795_SOURCE_SHEET = "RAAS hit validation"
ELIFE69795_SOURCE_COLUMNS = {
    "MTZ": "9mM MTZ",
    "Olmesartan": "9mM MTZ+Olmesartan",
    "Captopril": "9mM MTZ+Captopril",
    "Aliskiren": "9mM MTZ+Aliskiren",
    "Candesartan": "9mM MTZ+Candesartan",
}
DISC1_SOURCE_SHEET = "Fig. 1b"
DISC1_SOURCE_CONTEXT = "D2"
DISC1_SOURCE_ACTIONS = {
    "DMSO",
    "Rolipram",
    "NCGC00071192-01",
    "NCGC00067819-01",
    "NCGC00067541-01",
    "NCGC00071837-01",
    "NCGC00015985-01",
    "NCGC00065586-01",
    "NCGC00030627-01",
    "NCGC00024862-02",
    "NCGC00070528-01",
}
NATAGING_TDP43_SOURCE_SHEET = "Fig. 2"
NATAGING_TDP43_SURVIVAL_HEADER_ROW = 2
NATAGING_TDP43_SURVIVAL_REPLICATE_ROWS = (3, 4, 5)
NATAGING_TDP43_SURVIVAL_COLUMNS = {
    "DMSO": 56,
    "XL20": 57,
    "XL21": 58,
    "XL22": 59,
    "XL23": 60,
    "XL24": 61,
    "XL25": 62,
    "XL26": 63,
    "XL27": 64,
}


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _artifact(path: Path) -> dict[str, Any]:
    if not path.is_file() or path.is_symlink():
        raise FileNotFoundError(path)
    return {
        "path": path.relative_to(REPO_ROOT).as_posix(),
        "bytes": path.stat().st_size,
        "sha256": _sha256(path),
    }


def _canonical_hash(value: object) -> str:
    return hashlib.sha256(
        json.dumps(value, ensure_ascii=True, sort_keys=True, separators=(",", ":")).encode("utf-8")
    ).hexdigest()


def _repo_path(value: str) -> Path:
    path = (REPO_ROOT / value).resolve()
    if path != REPO_ROOT and REPO_ROOT not in path.parents:
        raise ValueError(f"Configured path escapes repository: {value}")
    return path


def _read_manifest(path: Path) -> list[dict[str, str]]:
    if not path.is_file():
        raise FileNotFoundError(path)
    with path.open("r", encoding="utf-8", newline="") as handle:
        reader = csv.DictReader(handle, delimiter="\t")
        if tuple(reader.fieldnames or ()) != MANIFEST_FIELDS:
            raise ValueError("Evidence materialization manifest columns drifted")
        rows = [{key: str(value or "").strip() for key, value in row.items()} for row in reader]
    if not rows:
        raise ValueError("Evidence materialization manifest is empty")
    return rows


def _write_tsv(path: Path, fields: Iterable[str], rows: Iterable[Mapping[str, Any]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(fields), delimiter="\t", extrasaction="ignore")
        writer.writeheader()
        writer.writerows({field: row.get(field, "") for field in fields} for row in rows)


def _load_contract() -> dict[str, Any]:
    payload = yaml.safe_load(CONTRACT_PATH.read_text(encoding="utf-8")) or {}
    contract = payload.get("retrospective_public_therapeutic_evidence")
    schema = json.loads(CONTRACT_SCHEMA_PATH.read_text(encoding="utf-8"))
    if not isinstance(contract, dict):
        raise TypeError("Retrospective public evidence contract must be a mapping")
    errors = sorted(Draft202012Validator(schema).iter_errors(contract), key=lambda error: list(error.absolute_path))
    if errors:
        detail = "; ".join(f"{list(error.absolute_path)}: {error.message}" for error in errors[:8])
        raise ValueError(f"Retrospective public evidence contract is invalid: {detail}")
    return contract


def _finite(value: str) -> bool:
    try:
        return math.isfinite(float(value))
    except (TypeError, ValueError):
        return False


def _figure3f_qc(row: Mapping[str, str]) -> dict[str, Any]:
    """Run two independent y-axis calibrations against the public raster figure."""

    locator = row["source_locator"]
    arm = locator.rsplit(":", 1)[-1]
    if arm == "vehicle":
        arm = "DMSO"
    if arm not in FIGURE3F_BAR_TOP_ROWS:
        raise ValueError(f"Unknown Figure 3F arm locator: {locator}")
    image = Image.open(_repo_path(row["source_artifact_path"])).convert("RGB")
    expected_top = FIGURE3F_BAR_TOP_ROWS[arm]
    center = FIGURE3F_PANEL_X_OFFSET + FIGURE3F_BAR_CENTERS[arm]
    fill = np.asarray(image)
    if fill.ndim != 3 or fill.shape[2] != 3:
        raise ValueError("Figure raster is not an RGB image")
    # The downloaded public PMC figure is a fixed 731 x 921 raster.  Exact-fill
    # pixels are used only to locate the bar; values are calibrated from axes.
    panel = fill[
        FIGURE3F_PANEL_Y_OFFSET:FIGURE3F_PANEL_Y_OFFSET + FIGURE3F_AXIS_ZERO_PX + 1,
        :,
    ]
    crop = panel[:230, max(0, center - 8):min(panel.shape[1], center + 9)]
    fill_mask = (
        (np.abs(crop[:, :, 0].astype(int) - 113) <= 8)
        & (np.abs(crop[:, :, 1].astype(int) - 123) <= 8)
        & (np.abs(crop[:, :, 2].astype(int) - 185) <= 8)
    )
    counts = fill_mask.sum(axis=1)
    candidate_rows = np.flatnonzero(counts >= 3)
    if candidate_rows.size == 0:
        raise ValueError(f"Figure bar fill not found for {arm}")
    detected_fill_top = int(candidate_rows[0])
    if abs(detected_fill_top - expected_top) > 6:
        raise ValueError(f"Figure bar top drift for {arm}: {detected_fill_top} vs {expected_top}")
    endpoint_mean = (FIGURE3F_AXIS_ZERO_PX - expected_top) * 400.0 / (FIGURE3F_AXIS_ZERO_PX - FIGURE3F_AXIS_TOP_PX)
    design = np.asarray(FIGURE3F_TICK_ROWS, dtype=float)
    values = np.asarray((0, 100, 200, 300, 400), dtype=float)
    slope, intercept = np.polyfit(design, values, deg=1)
    regression_mean = float(slope * expected_top + intercept)
    manifest_mean = float(row["arm_mean"])
    if abs(endpoint_mean - regression_mean) > 8.0 or abs(manifest_mean - regression_mean) > 10.0:
        raise ValueError(
            f"Figure calibration disagreement for {arm}: endpoint={endpoint_mean:.2f}, "
            f"regression={regression_mean:.2f}, manifest={manifest_mean:.2f}"
        )
    return {
        "arm": arm,
        "bar_center_px": center,
        "expected_bar_top_px": expected_top,
        "detected_fill_top_px": detected_fill_top,
        "endpoint_mean": endpoint_mean,
        "regression_mean": regression_mean,
        "manifest_mean": manifest_mean,
        "mean_disagreement": abs(endpoint_mean - regression_mean),
        "visible_error_bar_se": float(row["uncertainty"]),
    }


def _white_component_from_center(image: np.ndarray, center: tuple[int, int]) -> tuple[float, float, int]:
    """Recover one enclosed white replicate marker from a fixed public raster."""

    center_x, center_y = center
    radius = 24
    x0 = max(0, center_x - radius)
    y0 = max(0, center_y - radius)
    x1 = min(image.shape[1], center_x + radius + 1)
    y1 = min(image.shape[0], center_y + radius + 1)
    patch = image[y0:y1, x0:x1]
    mask = patch > 220
    seed_x = center_x - x0
    seed_y = center_y - y0
    if not (0 <= seed_y < mask.shape[0] and 0 <= seed_x < mask.shape[1] and mask[seed_y, seed_x]):
        raise ValueError(f"Expected replicate marker is not white at {center_x},{center_y}")
    stack = [(seed_y, seed_x)]
    visited = {(seed_y, seed_x)}
    while stack:
        y, x = stack.pop()
        for dy, dx in ((1, 0), (-1, 0), (0, 1), (0, -1)):
            next_y = y + dy
            next_x = x + dx
            if (
                0 <= next_y < mask.shape[0]
                and 0 <= next_x < mask.shape[1]
                and mask[next_y, next_x]
                and (next_y, next_x) not in visited
            ):
                visited.add((next_y, next_x))
                stack.append((next_y, next_x))
    coordinates = np.asarray(list(visited), dtype=float)
    return (
        float(coordinates[:, 1].mean() + x0),
        float(coordinates[:, 0].mean() + y0),
        len(visited),
    )


def _elife109518_fig5_qc(row: Mapping[str, str]) -> dict[str, Any]:
    """Validate digitized Figure 5H dots and direction-normalized GluSph summaries."""

    locator = row["source_locator"]
    prefix = "ELIFE109518:Figure5H:"
    if not locator.startswith(prefix):
        raise ValueError(f"Unexpected eLife 109518 Figure 5H locator: {locator}")
    group = locator[len(prefix):]
    expected_centers = ELIFE109518_FIG5_DOT_CENTERS.get(group)
    if expected_centers is None:
        raise ValueError(f"Unknown eLife 109518 Figure 5H group locator: {locator}")
    image = Image.open(_repo_path(row["source_artifact_path"])).convert("L")
    if image.size != (2205, 2700):
        raise ValueError(f"Unexpected eLife 109518 Figure 5 image size: {image.size}")
    tick_rows = np.asarray(ELIFE109518_FIG5_AXIS_TICK_ROWS, dtype=float)
    tick_values = np.asarray(ELIFE109518_FIG5_AXIS_TICK_VALUES, dtype=float)
    endpoint_slope = float((tick_values[-1] - tick_values[0]) / (tick_rows[-1] - tick_rows[0]))
    endpoint_intercept = float(tick_values[0] - endpoint_slope * tick_rows[0])
    regression_slope, regression_intercept = np.polyfit(tick_rows, tick_values, deg=1)
    gray = np.asarray(image, dtype=np.uint8)
    centroids: list[dict[str, Any]] = []
    raw_values_endpoint: list[float] = []
    raw_values_regression: list[float] = []
    for expected in expected_centers:
        centroid_x, centroid_y, area = _white_component_from_center(gray, expected)
        if area < 100 or area > 220:
            raise ValueError(f"Unexpected replicate marker area for {group}: {area}")
        if abs(centroid_x - expected[0]) > 3.0 or abs(centroid_y - expected[1]) > 3.0:
            raise ValueError(f"Replicate marker drift for {group}: {centroid_x:.2f},{centroid_y:.2f} vs {expected}")
        raw_values_endpoint.append(endpoint_slope * centroid_y + endpoint_intercept)
        raw_values_regression.append(float(regression_slope * centroid_y + regression_intercept))
        centroids.append({
            "expected_center_px": list(expected),
            "detected_center_px": [centroid_x, centroid_y],
            "white_component_area": area,
        })
    endpoint_values = np.asarray(raw_values_endpoint, dtype=float)
    regression_values = np.asarray(raw_values_regression, dtype=float)
    raw_mean_endpoint = float(endpoint_values.mean())
    raw_mean_regression = float(regression_values.mean())
    raw_sem_endpoint = float(endpoint_values.std(ddof=1) / math.sqrt(endpoint_values.size))
    raw_sem_regression = float(regression_values.std(ddof=1) / math.sqrt(regression_values.size))
    normalized_mean = -raw_mean_endpoint
    manifest_mean = float(row["arm_mean"])
    manifest_sem = float(row["uncertainty"])
    if abs(raw_mean_endpoint - raw_mean_regression) > 0.75:
        raise ValueError(
            f"Axis calibration disagreement for {group}: "
            f"endpoint={raw_mean_endpoint:.3f}, regression={raw_mean_regression:.3f}"
        )
    if abs(manifest_mean - normalized_mean) > 1.0:
        raise ValueError(
            f"Manifest mean disagrees with digitized direction-normalized value for {group}: "
            f"{manifest_mean:.3f} vs {normalized_mean:.3f}"
        )
    if abs(manifest_sem - raw_sem_endpoint) > 1.0:
        raise ValueError(
            f"Manifest SEM disagrees with digitized value for {group}: "
            f"{manifest_sem:.3f} vs {raw_sem_endpoint:.3f}"
        )
    if row["uncertainty_kind"] != "extracted_se":
        raise ValueError(f"Unexpected uncertainty kind for eLife 109518 group {group}")
    if row["extraction_qc_status"] != "dual_pass_calibration_verified":
        raise ValueError(f"Figure extraction QC not verified for eLife 109518 group {group}")
    return {
        "publication_uid": row["publication_uid"],
        "source_figure": "Figure 5H",
        "group": group,
        "axis_tick_rows_px": list(ELIFE109518_FIG5_AXIS_TICK_ROWS),
        "axis_tick_values": list(ELIFE109518_FIG5_AXIS_TICK_VALUES),
        "axis_calibration": {
            "endpoint_slope": endpoint_slope,
            "endpoint_intercept": endpoint_intercept,
            "regression_slope": float(regression_slope),
            "regression_intercept": float(regression_intercept),
            "mean_disagreement": abs(raw_mean_endpoint - raw_mean_regression),
        },
        "replicate_markers": centroids,
        "raw_values_endpoint": endpoint_values.tolist(),
        "raw_mean_endpoint": raw_mean_endpoint,
        "raw_sem_endpoint": raw_sem_endpoint,
        "raw_mean_regression": raw_mean_regression,
        "raw_sem_regression": raw_sem_regression,
        "direction_normalization": "negative_identity_lower_raw_GluSph_is_better",
        "manifest_mean": manifest_mean,
        "manifest_sem": manifest_sem,
        "individual_replicates_claimed": False,
    }


def _elife69795_source_data_qc(row: Mapping[str, str]) -> dict[str, Any]:
    """Recompute the exact arm summary from the public Figure 2 source workbook."""

    arm = row["source_locator"].rsplit(":", 1)[-1]
    source_header = ELIFE69795_SOURCE_COLUMNS.get(arm)
    if source_header is None:
        raise ValueError(f"Unknown eLife 69795 Figure 2 arm locator: {row['source_locator']}")
    workbook = openpyxl.load_workbook(
        _repo_path(row["source_artifact_path"]),
        read_only=True,
        data_only=True,
    )
    try:
        if ELIFE69795_SOURCE_SHEET not in workbook.sheetnames:
            raise ValueError(f"Missing source-data sheet: {ELIFE69795_SOURCE_SHEET}")
        worksheet = workbook[ELIFE69795_SOURCE_SHEET]
        headers = [worksheet.cell(1, column).value for column in range(1, worksheet.max_column + 1)]
        matching_columns = [
            column
            for column, header in enumerate(headers, start=1)
            if str(header or "").strip() == source_header
        ]
        if len(matching_columns) != 1:
            raise ValueError(f"Source-data arm {source_header!r} maps to {len(matching_columns)} columns")
        column = matching_columns[0]
        values = np.asarray(
            [
                float(worksheet.cell(source_row, column).value)
                for source_row in range(2, worksheet.max_row + 1)
                if isinstance(worksheet.cell(source_row, column).value, (int, float))
            ],
            dtype=float,
        )
    finally:
        workbook.close()
    if values.size < 2 or not np.isfinite(values).all():
        raise ValueError(f"Invalid public replicate values for {source_header}")
    observed_mean = float(np.mean(values))
    observed_sd = float(np.std(values, ddof=1))
    manifest_n = int(row["biological_replicate_count"])
    manifest_mean = float(row["arm_mean"])
    manifest_sd = float(row["uncertainty"])
    if manifest_n != int(values.size):
        raise ValueError(f"Replicate count drift for {source_header}: {manifest_n} vs {values.size}")
    if not math.isclose(manifest_mean, observed_mean, rel_tol=1e-12, abs_tol=1e-12):
        raise ValueError(f"Mean drift for {source_header}: {manifest_mean} vs {observed_mean}")
    if not math.isclose(manifest_sd, observed_sd, rel_tol=1e-12, abs_tol=1e-12):
        raise ValueError(f"SD drift for {source_header}: {manifest_sd} vs {observed_sd}")
    if row["uncertainty_kind"] != "extracted_sd":
        raise ValueError(f"Unexpected uncertainty kind for {source_header}: {row['uncertainty_kind']}")
    if row["extraction_qc_status"] != "source_data_replicate_summary_verified":
        raise ValueError(f"Source-data extraction QC not verified for {source_header}")
    return {
        "publication_uid": row["publication_uid"],
        "source_sheet": ELIFE69795_SOURCE_SHEET,
        "source_header": source_header,
        "replicate_count": int(values.size),
        "recomputed_mean": observed_mean,
        "recomputed_sd": observed_sd,
        "manifest_mean": manifest_mean,
        "manifest_sd": manifest_sd,
    }


def _disc1_source_data_qc(row: Mapping[str, str]) -> dict[str, Any]:
    """Recompute the prespecified D2 Figure 1b arm summary from source data.

    The source article identifies the three values per arm as independent
    differentiation cultures.  The D2 patient line is the first patient line
    in the source table and is fixed before candidate qualification; the D3
    line is not a second split group from the same experimental campaign.
    """

    action = row["source_locator"].rsplit(":", 1)[-1]
    if action not in DISC1_SOURCE_ACTIONS:
        raise ValueError(f"Unknown DISC1 Figure 1b arm locator: {row['source_locator']}")
    workbook = openpyxl.load_workbook(
        _repo_path(row["source_artifact_path"]),
        read_only=True,
        data_only=True,
    )
    try:
        if DISC1_SOURCE_SHEET not in workbook.sheetnames:
            raise ValueError(f"Missing source-data sheet: {DISC1_SOURCE_SHEET}")
        rows = list(workbook[DISC1_SOURCE_SHEET].iter_rows(values_only=True))
    finally:
        workbook.close()
    if len(rows) < 5:
        raise ValueError("DISC1 Figure 1b source data has fewer than three replicate rows")
    context_columns = [
        index for index, value in enumerate(rows[0]) if str(value or "").strip() == DISC1_SOURCE_CONTEXT
    ]
    if len(context_columns) != 1:
        raise ValueError(
            f"DISC1 source context {DISC1_SOURCE_CONTEXT!r} maps to {len(context_columns)} columns"
        )
    context_start = context_columns[0]
    next_context_starts = [
        index
        for index, value in enumerate(rows[0])
        if index > context_start and str(value or "").strip()
    ]
    context_end = next_context_starts[0] if next_context_starts else len(rows[1])
    action_columns = [
        index
        for index in range(context_start, context_end)
        if str(rows[1][index] or "").strip() == action
    ]
    if len(action_columns) != 1:
        raise ValueError(
            f"DISC1 source action {action!r} maps to {len(action_columns)} columns "
            f"within context {DISC1_SOURCE_CONTEXT!r}"
        )
    column = action_columns[0]
    values = np.asarray(
        [
            float(source_row[column])
            for source_row in rows[2:]
            if column < len(source_row) and isinstance(source_row[column], (int, float))
        ],
        dtype=float,
    )
    if values.size != 3 or not np.isfinite(values).all():
        raise ValueError(f"Invalid independent-differentiation values for DISC1 action {action}")
    observed_mean = float(np.mean(values))
    observed_se = float(np.std(values, ddof=1) / math.sqrt(values.size))
    manifest_n = int(row["biological_replicate_count"])
    manifest_mean = float(row["arm_mean"])
    manifest_se = float(row["uncertainty"])
    if manifest_n != int(values.size):
        raise ValueError(f"Replicate count drift for DISC1 action {action}: {manifest_n} vs {values.size}")
    if not math.isclose(manifest_mean, observed_mean, rel_tol=1e-12, abs_tol=1e-12):
        raise ValueError(f"Mean drift for DISC1 action {action}: {manifest_mean} vs {observed_mean}")
    if not math.isclose(manifest_se, observed_se, rel_tol=1e-12, abs_tol=1e-12):
        raise ValueError(f"SEM drift for DISC1 action {action}: {manifest_se} vs {observed_se}")
    if row["uncertainty_kind"] != "extracted_se":
        raise ValueError(f"Unexpected uncertainty kind for DISC1 action {action}: {row['uncertainty_kind']}")
    if row["extraction_qc_status"] != "source_data_independent_differentiation_verified":
        raise ValueError(f"DISC1 source-data extraction QC not verified for {action}")
    return {
        "publication_uid": row["publication_uid"],
        "source_sheet": DISC1_SOURCE_SHEET,
        "source_context": DISC1_SOURCE_CONTEXT,
        "source_action": action,
        "replicate_count": int(values.size),
        "recomputed_mean": observed_mean,
        "recomputed_se": observed_se,
        "manifest_mean": manifest_mean,
        "manifest_se": manifest_se,
    }


def _nataging_tdp43_source_data_qc(row: Mapping[str, str]) -> dict[str, Any]:
    """Recompute the prespecified primary-neuron survival summaries from source data.

    The source workbook block is fixed by its Figure 2 treatment headers.  GFP
    is a transduction reference outside the therapeutic action set, so it is
    deliberately excluded rather than being treated as an intervention arm.
    """

    action = row["source_locator"].rsplit(":", 1)[-1]
    if action not in NATAGING_TDP43_SURVIVAL_COLUMNS:
        raise ValueError(f"Unknown TDP-43 Figure 2 survival arm locator: {row['source_locator']}")
    workbook = openpyxl.load_workbook(
        _repo_path(row["source_artifact_path"]),
        read_only=True,
        data_only=True,
    )
    try:
        if NATAGING_TDP43_SOURCE_SHEET not in workbook.sheetnames:
            raise ValueError(f"Missing source-data sheet: {NATAGING_TDP43_SOURCE_SHEET}")
        rows = list(workbook[NATAGING_TDP43_SOURCE_SHEET].iter_rows(values_only=True))
    finally:
        workbook.close()
    column = NATAGING_TDP43_SURVIVAL_COLUMNS[action]
    if len(rows) <= max(NATAGING_TDP43_SURVIVAL_REPLICATE_ROWS):
        raise ValueError("TDP-43 Figure 2 source data has fewer than three replicate rows")
    if column >= len(rows[NATAGING_TDP43_SURVIVAL_HEADER_ROW]):
        raise ValueError(f"TDP-43 Figure 2 column is missing for action {action}")
    if str(rows[NATAGING_TDP43_SURVIVAL_HEADER_ROW][column] or "").strip() != action:
        raise ValueError(f"TDP-43 Figure 2 header drift for action {action}")
    values = np.asarray(
        [float(rows[index][column]) for index in NATAGING_TDP43_SURVIVAL_REPLICATE_ROWS],
        dtype=float,
    )
    if values.size != 3 or not np.isfinite(values).all():
        raise ValueError(f"Invalid independent-neuronal-culture values for TDP-43 action {action}")
    observed_mean = float(np.mean(values))
    observed_se = float(np.std(values, ddof=1) / math.sqrt(values.size))
    manifest_n = int(row["biological_replicate_count"])
    manifest_mean = float(row["arm_mean"])
    manifest_se = float(row["uncertainty"])
    if manifest_n != int(values.size):
        raise ValueError(f"Replicate count drift for TDP-43 action {action}: {manifest_n} vs {values.size}")
    if not math.isclose(manifest_mean, observed_mean, rel_tol=1e-12, abs_tol=1e-12):
        raise ValueError(f"Mean drift for TDP-43 action {action}: {manifest_mean} vs {observed_mean}")
    if not math.isclose(manifest_se, observed_se, rel_tol=1e-12, abs_tol=1e-12):
        raise ValueError(f"SEM drift for TDP-43 action {action}: {manifest_se} vs {observed_se}")
    if row["uncertainty_kind"] != "extracted_se":
        raise ValueError(f"Unexpected uncertainty kind for TDP-43 action {action}: {row['uncertainty_kind']}")
    if row["extraction_qc_status"] != "source_data_independent_neuronal_culture_verified":
        raise ValueError(f"TDP-43 source-data extraction QC not verified for {action}")
    return {
        "publication_uid": row["publication_uid"],
        "source_sheet": NATAGING_TDP43_SOURCE_SHEET,
        "source_block": "primary_cortical_neuron_survival",
        "source_action": action,
        "replicate_count": int(values.size),
        "recomputed_mean": observed_mean,
        "recomputed_se": observed_se,
        "manifest_mean": manifest_mean,
        "manifest_se": manifest_se,
    }


def _validate_manifest_rows(rows: list[dict[str, str]]) -> tuple[list[str], list[str]]:
    blockers: list[str] = []
    accepted: list[dict[str, str]] = []
    record_uids: set[str] = set()
    for index, row in enumerate(rows, start=2):
        prefix = f"row_{index}"
        record_uid = row["record_uid"]
        if not record_uid or record_uid in record_uids:
            blockers.append(f"{prefix}:duplicate_or_missing_record_uid")
        record_uids.add(record_uid)
        if row["record_status"] not in {"accepted", "rejected"}:
            blockers.append(f"{prefix}:invalid_record_status")
        if row["public_access_status"] != "public":
            blockers.append(f"{prefix}:non_public_source")
        source_path = _repo_path(row["source_artifact_path"]) if row["source_artifact_path"] else None
        if source_path is None or not source_path.is_file() or source_path.is_symlink():
            blockers.append(f"{prefix}:source_artifact_missing")
        else:
            actual_hash = _sha256(source_path)
            if row["source_content_hash"] != actual_hash:
                blockers.append(f"{prefix}:source_content_hash_mismatch")
        if len(row["source_content_hash"]) != 64 or not all(c in "0123456789abcdefABCDEF" for c in row["source_content_hash"]):
            blockers.append(f"{prefix}:invalid_source_content_hash")
        if row["record_status"] == "rejected":
            if not row["rejection_reason"]:
                blockers.append(f"{prefix}:rejection_reason_missing")
            continue
        accepted.append(row)
        if row["evidence_grade"] not in FORMAL_GRADES:
            blockers.append(f"{prefix}:accepted_non_formal_grade")
        for field in EVIDENCE_FIELDS:
            if field == "decision_domain_classification":
                continue
            if field == "source_content_hash":
                continue
            if not row.get(field, "") and field not in {"replicate_uid", "numeric_outcome", "split_role"}:
                blockers.append(f"{prefix}:missing_{field}")
        if row["evidence_grade"] in {"B", "C"}:
            if not _finite(row["arm_mean"]):
                blockers.append(f"{prefix}:invalid_arm_mean")
            try:
                if int(row["biological_replicate_count"]) < 2:
                    blockers.append(f"{prefix}:replicate_count_below_two")
            except ValueError:
                blockers.append(f"{prefix}:invalid_replicate_count")
            if not _finite(row["uncertainty"]) or float(row["uncertainty"]) <= 0:
                blockers.append(f"{prefix}:invalid_uncertainty")
        if row["evidence_grade"] == "C":
            extraction_path = source_path
            if not row["extraction_method"] or not row["extraction_qc_status"]:
                blockers.append(f"{prefix}:extraction_qc_metadata_missing")
            if len(row["extraction_artifact_hash"]) != 64 or not all(c in "0123456789abcdefABCDEF" for c in row["extraction_artifact_hash"]):
                blockers.append(f"{prefix}:invalid_extraction_artifact_hash")
            elif extraction_path is not None and row["extraction_artifact_hash"] != _sha256(extraction_path):
                blockers.append(f"{prefix}:extraction_artifact_hash_mismatch")
            if row["source_locator"].startswith("PMC13099497:Figure3F:"):
                try:
                    qc = _figure3f_qc(row)
                    if row["extraction_qc_status"] != "dual_pass_calibration_verified":
                        blockers.append(f"{prefix}:figure_extraction_qc_not_verified")
                    row["_figure_qc"] = json.dumps(qc, ensure_ascii=True, sort_keys=True)
                except (OSError, ValueError) as exc:
                    blockers.append(f"{prefix}:figure_extraction_qc:{exc}")
            if row["source_locator"].startswith("ELIFE69795:Figure2:SourceData1:"):
                try:
                    qc = _elife69795_source_data_qc(row)
                    row["_source_data_qc"] = json.dumps(qc, ensure_ascii=True, sort_keys=True)
                except (OSError, ValueError) as exc:
                    blockers.append(f"{prefix}:source_data_extraction_qc:{exc}")
            if row["source_locator"].startswith("PMC7930023:Figure1B:SourceData:D2:"):
                try:
                    qc = _disc1_source_data_qc(row)
                    row["_source_data_qc"] = json.dumps(qc, ensure_ascii=True, sort_keys=True)
                except (OSError, ValueError) as exc:
                    blockers.append(f"{prefix}:source_data_extraction_qc:{exc}")
            if row["source_locator"].startswith("NATAGING2026:Figure2:SourceData:survival:"):
                try:
                    qc = _nataging_tdp43_source_data_qc(row)
                    row["_source_data_qc"] = json.dumps(qc, ensure_ascii=True, sort_keys=True)
                except (OSError, ValueError) as exc:
                    blockers.append(f"{prefix}:source_data_extraction_qc:{exc}")
            if row["source_locator"].startswith("ELIFE109518:Figure5H:"):
                try:
                    qc = _elife109518_fig5_qc(row)
                    row["_figure_qc"] = json.dumps(qc, ensure_ascii=True, sort_keys=True)
                except (OSError, ValueError) as exc:
                    blockers.append(f"{prefix}:figure_extraction_qc:{exc}")
    return sorted(set(blockers)), accepted


def _summary_arm(row: Mapping[str, str]):
    n = int(row["biological_replicate_count"])
    mean = float(row["arm_mean"])
    uncertainty = float(row["uncertainty"])
    if row["uncertainty_kind"] in SUMMARY_KINDS:
        return build_summary_arm(mean=mean, n=n, sem=uncertainty)
    if row["uncertainty_kind"] in SD_KINDS:
        return build_summary_arm(mean=mean, n=n, sd=uncertainty)
    raise ValueError(f"Unsupported summary uncertainty kind: {row['uncertainty_kind']}")


def _classify_candidate_set(rows: list[dict[str, str]]) -> dict[str, Any]:
    controls = {row["common_control_uid"] for row in rows if row["common_control_uid"]}
    if len(controls) != 1:
        raise ValueError("A candidate set must have exactly one common control")
    control_uid = next(iter(controls))
    control_rows = [row for row in rows if row["structured_action_uid"] == control_uid]
    action_rows = [row for row in rows if row["structured_action_uid"] != control_uid]
    if len(control_rows) != 1 or len(action_rows) < 2:
        raise ValueError("A candidate set requires one control and at least two actions")
    control = _summary_arm(control_rows[0])
    rng = np.random.default_rng(BOOTSTRAP_SEED)
    control_draws = rng.normal(control.mean, control.sd / math.sqrt(control.n), size=BOOTSTRAP_RESAMPLES)
    action_draws = []
    effects = []
    for row in action_rows:
        arm = _summary_arm(row)
        draws = rng.normal(arm.mean, arm.sd / math.sqrt(arm.n), size=BOOTSTRAP_RESAMPLES)
        action_draws.append(draws)
        effects.append(draws - control_draws)
    effect_matrix = np.asarray(effects, dtype=float)
    point_effects = np.asarray([float(row["arm_mean"]) - control.mean for row in action_rows], dtype=float)
    winners = np.argmax(effect_matrix, axis=0)
    top_probabilities = np.asarray([(winners == index).mean() for index in range(len(action_rows))])
    benefit_probabilities = np.asarray([(effect_matrix[index] > 0.0).mean() for index in range(len(action_rows))])
    winner_index = int(np.argmax(point_effects))
    ranked = np.sort(point_effects)[::-1]
    margin = float(ranked[0] - ranked[1])
    winner_top_probability = float(top_probabilities[winner_index])
    winner_benefit_probability = float(benefit_probabilities[winner_index])
    qualification = "ambiguous"
    if float(np.max(np.abs(point_effects))) <= 0.10:
        qualification = "near_null"
    elif winner_top_probability >= 0.95 and winner_benefit_probability >= 0.95 and margin >= 0.10:
        qualification = "decision_sensitive"
    elif winner_top_probability >= 0.80 and winner_benefit_probability >= 0.80 and margin >= 0.10:
        qualification = "weakly_decision_sensitive"
    return {
        "decision_domain_classification": qualification,
        "control_uid": control_uid,
        "action_uids": [row["structured_action_uid"] for row in action_rows],
        "point_effects": {row["structured_action_uid"]: float(value) for row, value in zip(action_rows, point_effects)},
        "top_probabilities": {row["structured_action_uid"]: float(value) for row, value in zip(action_rows, top_probabilities)},
        "benefit_probabilities": {row["structured_action_uid"]: float(value) for row, value in zip(action_rows, benefit_probabilities)},
        "winner_uid": action_rows[winner_index]["structured_action_uid"],
        "winner_top_probability": winner_top_probability,
        "winner_benefit_probability": winner_benefit_probability,
        "point_estimate_margin": margin,
        "bootstrap_resamples": BOOTSTRAP_RESAMPLES,
        "bootstrap_seed": BOOTSTRAP_SEED,
        "statistic_family": "parametric_normal_summary_bootstrap",
    }


def _candidate_groups(rows: list[dict[str, str]]) -> list[dict[str, str]]:
    groups: dict[str, list[dict[str, str]]] = {}
    for row in rows:
        groups.setdefault(row["candidate_set_uid"], []).append(row)
    for candidate_set_uid, members in groups.items():
        classification = _classify_candidate_set(members)
        for row in members:
            row["decision_domain_classification"] = classification["decision_domain_classification"]
    return [
        {"candidate_set_uid": key, "members": members, "classification": _classify_candidate_set(members)}
        for key, members in sorted(groups.items())
    ]


def build(*, write_output: bool = True) -> dict[str, Any]:
    contract = _load_contract()
    rows = _read_manifest(MANIFEST_PATH)
    manifest_blockers, accepted_rows = _validate_manifest_rows(rows)
    rejected_rows = [row for row in rows if row["record_status"] == "rejected"]
    if not manifest_blockers:
        accepted_blockers = validate_evidence_rows(accepted_rows)
        manifest_blockers.extend(f"accepted_evidence:{item}" for item in accepted_blockers)
    candidate_summaries: list[dict[str, Any]] = []
    if not manifest_blockers and accepted_rows:
        candidate_summaries = _candidate_groups(accepted_rows)
    candidates, candidate_blockers = qualify_candidate_sets(accepted_rows) if not manifest_blockers else ([], [])
    legal = [row for row in candidates if row["legal"]]
    qualified = [row for row in candidates if row["qualified"]]
    qualified_candidate_uids = {row["candidate_set_uid"] for row in qualified}
    qualified_evidence_rows = [
        row for row in accepted_rows if row["candidate_set_uid"] in qualified_candidate_uids
    ]
    independent_publications = len({row["publication_uid"] for row in accepted_rows})
    independent_contexts = len({row["biological_context_uid"] for row in accepted_rows})
    independent_groups = len({(row["publication_uid"], row["biological_context_uid"]) for row in accepted_rows})
    independent_decision_sensitive_groups = len(
        {
            (row["publication_uid"], row["biological_context_uid"])
            for row in qualified_evidence_rows
        }
    )
    replay_roles_feasible = bool(
        independent_decision_sensitive_groups >= MIN_INDEPENDENT_GROUPS_FOR_REPLAY_ROLES
    )
    if manifest_blockers:
        status = "blocked_materialization_input"
    elif not accepted_rows:
        status = "blocked_insufficient_public_quantitative_evidence"
    elif not legal:
        status = "public_quantitative_evidence_materialized_no_legal_candidate_sets"
    elif not qualified:
        status = "qualified_public_candidates_not_decision_sensitive"
    elif not replay_roles_feasible:
        status = "qualified_public_candidates_but_split_insufficient"
    else:
        status = "public_retrospective_candidate_domain_qualified"
    audit = {
        "schema_version": "neuronsifter-retrospective-public-evidence-materialization-audit-v1",
        "contract_version": contract["contract_version"],
        "status": status,
        "created_at": datetime.now(timezone.utc).isoformat(),
        "contract_digest": _canonical_hash(contract),
        "contract_binding": _artifact(CONTRACT_PATH),
        "contract_schema_binding": _artifact(CONTRACT_SCHEMA_PATH),
        "materialization_schema_binding": _artifact(MATERIALIZATION_SCHEMA_PATH),
        "implementation_binding": _artifact(BUILDER_PATH),
        "input_manifest_binding": _artifact(MANIFEST_PATH),
        "source_artifact_bindings": {
            row["source_artifact_path"]: _artifact(_repo_path(row["source_artifact_path"]))
            for row in rows
            if row["source_artifact_path"]
        },
        "accepted_record_count": len(accepted_rows),
        "rejected_record_count": len(rejected_rows),
        "formal_evidence_grade_counts": {
            grade: sum(row["evidence_grade"] == grade for row in accepted_rows)
            for grade in sorted(FORMAL_GRADES)
        },
        "candidate_set_count": len(candidates),
        "legal_candidate_set_count": len(legal),
        "decision_sensitive_candidate_set_count": len(qualified),
        "independent_publication_count": independent_publications,
        "independent_context_count": independent_contexts,
        "independent_replay_group_count": independent_groups,
        "independent_decision_sensitive_group_count": independent_decision_sensitive_groups,
        "minimum_independent_replay_groups": MIN_INDEPENDENT_GROUPS_FOR_REPLAY_ROLES,
        "replay_roles_feasible": replay_roles_feasible,
        "candidate_statistics": candidate_summaries,
        "figure_extraction_qc": [
            json.loads(row["_figure_qc"])
            for row in accepted_rows
            if row.get("_figure_qc")
        ],
        "source_data_extraction_qc": [
            json.loads(row["_source_data_qc"])
            for row in accepted_rows
            if row.get("_source_data_qc")
        ],
        "candidate_qualification": candidates,
        "candidate_blockers": sorted(set(candidate_blockers)),
        "manifest_blockers": sorted(set(manifest_blockers)),
        "objective_selection_outcome_blind": contract["admission"]["objective_selection_outcome_blind"],
        "outcome_values_used_for_source_priority": False,
        "model_predictions_inspected": False,
        "protected_holdout_access": 0,
        "publication_claim_eligible": False,
        "world_action_authorized": False,
        "agent_in_twin_authorized": False,
        "coevolution_authorized": False,
        "output_bindings": {
            "evidence_registry": {"path": EVIDENCE_OUTPUT_PATH.relative_to(REPO_ROOT).as_posix()},
            "rejection_catalog": {"path": REJECTION_OUTPUT_PATH.relative_to(REPO_ROOT).as_posix()},
        },
        "next_route": (
            "materialize_independent_public_contexts_and_sealed_snapshots"
            if qualified and not replay_roles_feasible
            else "formal_digital_twin_qualification"
            if replay_roles_feasible
            else "bounded_public_quantitative_evidence_materialization"
        ),
    }
    evidence_rows = [{field: row.get(field, "") for field in EVIDENCE_FIELDS} for row in accepted_rows]
    rejection_rows = [
        {field: row.get(field, "") for field in REJECTION_FIELDS}
        for row in rejected_rows
    ]
    if write_output:
        _write_tsv(EVIDENCE_OUTPUT_PATH, EVIDENCE_FIELDS, evidence_rows)
        _write_tsv(REJECTION_OUTPUT_PATH, REJECTION_FIELDS, rejection_rows)
        AUDIT_OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
        AUDIT_OUTPUT_PATH.write_text(json.dumps(audit, ensure_ascii=True, indent=2, sort_keys=True) + "\n", encoding="utf-8")
    return audit


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--no-write", action="store_true")
    args = parser.parse_args()
    print(json.dumps(build(write_output=not args.no_write), ensure_ascii=True, indent=2, sort_keys=True))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
